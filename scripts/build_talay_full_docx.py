#!/usr/bin/env python3
"""Сборка полного отчёта ЖК «Талай» в Word DOCX."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path("/workspace")
PARTS_DIR = ROOT / "docs/deliverables/full_report_parts"
OUT = ROOT / "docs/deliverables/word/ЖК_Талай_полный_отчёт_IMPRO_СД.docx"

PART_FILES = [
    "01_03_резюме_исследование_рынок.md",
    "04_06_проект_конкуренты_продукт.md",
    "07_10_аудитория_маркетинг_продажи_агенты.md",
    "11_13_swot_риски_возможности.md",
    "14_16_инициативы_приоритеты_дорожная_карта.md",
    "17_19_сопровождение_модель_подрядчики.md",
]


def set_run_font(run, *, bold=False, italic=False, size=11, name="Times New Roman"):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)


def add_formatted_text(paragraph, text: str, *, bold=False, italic=False, size=11):
    """Поддержка **жирного** и *курсива* в одной строке."""
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, bold=True, italic=italic, size=size)
        elif part.startswith("*") and part.endswith("*") and not part.startswith("**"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, bold=bold, italic=True, size=size)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, bold=bold, italic=italic, size=size)


def style_paragraph(p, *, space_after=4, space_before=0, first_line=True):
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.1
    if first_line:
        pf.first_line_indent = Cm(0.9)


def clean_md(text: str) -> str:
    text = text.replace("\r\n", "\n")
    # убрать лишние якоря путей в тексте отчёта для основателя
    text = text.replace("BD PMO", "внутренний рабочий ряд")
    text = re.sub(r"`[^`]+`", lambda m: m.group(0).strip("`"), text)
    return text


def parse_table_block(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(re.match(r"^:?-+:?$", c.replace(" ", "")) for c in cells):
            continue
        rows.append([re.sub(r"\*\*", "", c) for c in cells])
    return rows


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            val = row[j] if j < len(row) else ""
            add_formatted_text(p, val, bold=(i == 0), size=8)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            for run in p.runs:
                run.font.name = "Times New Roman"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                run.font.size = Pt(8)
    doc.add_paragraph()


def add_heading_custom(doc: Document, text: str, level: int):
    text = re.sub(r"\*\*", "", text).strip()
    # убрать ведущие # уже сняты
    h = doc.add_heading(text, level=min(level, 3))
    for run in h.runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        if level == 1:
            run.font.size = Pt(16)
        elif level == 2:
            run.font.size = Pt(14)
        else:
            run.font.size = Pt(12)


def render_markdown(doc: Document, md: str):
    md = clean_md(md)
    lines = md.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # horizontal rule
        if re.match(r"^-{3,}$", stripped) or re.match(r"^\*{3,}$", stripped):
            i += 1
            continue

        # headings
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            add_heading_custom(doc, m.group(2), level)
            i += 1
            continue

        # table
        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, parse_table_block(block))
            continue

        # bullet
        if re.match(r"^[-*•]\s+", stripped):
            text = re.sub(r"^[-*•]\s+", "", stripped)
            p = doc.add_paragraph(style="List Bullet")
            add_formatted_text(p, text, size=11)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(3)
            i += 1
            continue

        # numbered
        if re.match(r"^\d+[.)]\s+", stripped):
            text = re.sub(r"^\d+[.)]\s+", "", stripped)
            p = doc.add_paragraph(style="List Number")
            add_formatted_text(p, text, size=11)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(3)
            i += 1
            continue

        # blockquote
        if stripped.startswith(">"):
            text = re.sub(r"^>\s?", "", stripped)
            p = doc.add_paragraph()
            add_formatted_text(p, text, italic=True, size=11)
            style_paragraph(p, first_line=False, space_before=6, space_after=6)
            p.paragraph_format.left_indent = Cm(1)
            i += 1
            continue

        # normal paragraph — склеиваем продолжения
        para_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (
                not nxt
                or nxt.startswith("#")
                or nxt.startswith("|")
                or re.match(r"^[-*•]\s+", nxt)
                or re.match(r"^\d+[.)]\s+", nxt)
                or nxt.startswith(">")
                or re.match(r"^-{3,}$", nxt)
            ):
                break
            para_lines.append(nxt)
            i += 1
        text = " ".join(para_lines)
        p = doc.add_paragraph()
        add_formatted_text(p, text, size=10.5)
        style_paragraph(p)


def add_title_page(doc: Document):
    for _ in range(3):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("IMPRO × Сибирь Девелопмент")
    set_run_font(run, size=14, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ЖК «ТАЛАЙ»")
    set_run_font(run, size=28, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "Полный аналитический отчёт\n"
        "для обсуждения с основателем и директором по маркетингу"
    )
    set_run_font(run, size=14)

    doc.add_paragraph()

    meta = [
        "Горно-Алтайск · запуск продаж сентябрь–декабрь 2026",
        "Цель периода: не менее 300 млн ₽ по зарегистрированным ДДУ",
        "Версия 1.0 · 30 июля 2026",
        "Документ подготовлен на основе открытых данных, меморандумов проекта",
        "и исследовательских слоёв IMPRO. Не является офертой и не гарантирует сделки.",
    ]
    for line in meta:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        set_run_font(run, size=11)
        p.paragraph_format.space_after = Pt(4)

    doc.add_page_break()


def add_toc_outline(doc: Document):
    add_heading_custom(doc, "Содержание", 1)
    items = [
        "1. Краткое резюме для руководства",
        "2. Описание исследования",
        "3. Анализ рынка",
        "4. Анализ проекта «Талай»",
        "5. Анализ конкурентов",
        "6. Анализ продукта",
        "7. Анализ целевой аудитории",
        "8. Анализ маркетинга",
        "9. Анализ продаж",
        "10. Анализ агентского канала",
        "11. СВОТ-анализ",
        "12. Основные риски",
        "13. Возможности роста",
        "14. Стратегические инициативы",
        "15. Приоритеты внедрения",
        "16. Дорожная карта",
        "17. Почему инициативы требуют постоянного сопровождения",
        "18. Предлагаемая модель взаимодействия",
        "19. Подрядчики и их контроль",
        "Заключение",
    ]
    for item in items:
        p = doc.add_paragraph()
        add_formatted_text(p, item, size=12)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    doc.add_page_break()


def add_method_note(doc: Document):
    add_heading_custom(doc, "Как читать этот документ", 1)
    notes = [
        "Документ собран по полочкам: от рынка и проекта к инициативам и модели взаимодействия. Его можно читать целиком или по главам.",
        "Метки внутри текста: «Факт» — подтверждено источником; «Вывод IMPRO» — наша интерпретация; «Гипотеза» — проверить на первых бронях; «Сценарий» — модель, не достигнутый рынок.",
        "Цифры по продукту и плану продаж взяты из меморандумов «Талай». Рыночные якоря — из Росстата, ЦБ, Домклик и открытых заявлений властей Республики Алтай на дату среза.",
        "Документ написан простым русским языком специально для стола основателя: без чужого профессионального жаргона и без «продажи пакета» вместо разбора дела.",
    ]
    for n in notes:
        p = doc.add_paragraph()
        add_formatted_text(p, n, size=11)
        style_paragraph(p)
    doc.add_page_break()


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(1.7)
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    add_title_page(doc)
    add_toc_outline(doc)
    add_method_note(doc)

    for idx, name in enumerate(PART_FILES):
        path = PARTS_DIR / name
        if not path.exists():
            raise FileNotFoundError(path)
        md = path.read_text(encoding="utf-8")
        render_markdown(doc, md)
        # без принудительного разрыва: главы идут потоком, экономим пустые хвосты страниц
        if idx < len(PART_FILES) - 1:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)

    # нижний колонтитул
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("ЖК «Талай» · полный отчёт IMPRO × Сибирь Девелопмент · конфиденциально")
    set_run_font(run, size=9, italic=True)

    doc.save(OUT)
    print(f"Saved: {OUT}")
    print(f"Size: {OUT.stat().st_size / 1024:.1f} KB")


if __name__ == "__main__":
    main()
