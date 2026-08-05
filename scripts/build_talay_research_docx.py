#!/usr/bin/env python3
"""Сборка научного исследования рынка Алтая / кейс «Талай» в Word."""

from __future__ import annotations

import importlib.util
from pathlib import Path

ROOT = Path("/workspace")
OUT = ROOT / "docs/deliverables/word/Исследование_рынок_жилья_Горно_Алтайск_кейс_Талай.docx"

# переиспользуем рендер из предыдущего сборщика
spec = importlib.util.spec_from_file_location(
    "build_talay", ROOT / "scripts/build_talay_full_docx.py"
)
mod = importlib.util.module_from_spec(spec)
spec.loader.exec_module(mod)

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_run_font(run, *, bold=False, italic=False, size=12, name="Times New Roman"):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)


def add_title_page(doc: Document):
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ПРИКЛАДНОЕ ИССЛЕДОВАНИЕ")
    set_run_font(r, size=12, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "Рынок многоквартирного жилья агломерации\n"
        "Горно-Алтайск — Майма (Республика Алтай):\n"
        "структура предложения, платёжеспособный спрос\n"
        "и кейс городского жилого комплекса «Талай»"
    )
    set_run_font(r, size=16, bold=True)

    doc.add_paragraph()
    meta = [
        "Объект исследования: рынок первичного многоквартирного жилья агломерации",
        "Предмет: ценовые полки, конкуренция, спрос и позиционирование городского продукта",
        "Кейс приложения: ЖК «Талай» (Горно-Алтайск)",
        "Подготовлено: IMPRO Group · август 2026",
        "Жанр: описательно-аналитическое кабинетное исследование по открытым данным",
        "и проектным материалам объекта. Не является офертой и рекламой.",
    ]
    for line in meta:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        set_run_font(r, size=11)
        p.paragraph_format.space_after = Pt(4)

    doc.add_page_break()


def add_toc(doc: Document):
    p = doc.add_paragraph()
    r = p.add_run("Оглавление")
    set_run_font(r, size=14, bold=True)
    items = [
        "Аннотация",
        "1. Введение",
        "2. Методология и источники",
        "3. Макросреда Республики Алтай и агломерации",
        "4. Горно-Алтайск как город-ворота",
        "5. Предложение на рынке многоквартирного жилья",
        "6. Конкурентный анализ",
        "7. Спрос и платёжеспособность",
        "8. Кейс: жилой комплекс «Талай»",
        "9. Расчётные сценарии доходности городской квартиры",
        "10. Канал посредников и организация продаж",
        "11. Результаты исследования",
        "12. Обсуждение и рабочие гипотезы",
        "13. Ограничения исследования",
        "14. Заключение",
        "Список источников",
        "Приложение. Добор: посредники, конкуренты, общественные помещения",
    ]
    for it in items:
        p = doc.add_paragraph()
        r = p.add_run(it)
        set_run_font(r, size=12)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.5
    doc.add_page_break()


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)

    # академическая вёрстка поверх общего рендера
    def style_paragraph(p, *, space_after=8, space_before=0, first_line=True):
        pf = p.paragraph_format
        pf.space_after = Pt(space_after)
        pf.space_before = Pt(space_before)
        pf.line_spacing = 1.5
        if first_line:
            pf.first_line_indent = Cm(1.25)

    mod.style_paragraph = style_paragraph

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)  # академический левый
    section.right_margin = Cm(1.5)
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.first_line_indent = Cm(1.25)

    add_title_page(doc)
    add_toc(doc)

    parts = [
        ROOT / "docs/deliverables/research/01_титул_введение_метод_макро.md",
        ROOT / "docs/deliverables/research/02_город_предложение_спрос_конкуренты.md",
        ROOT / "docs/deliverables/research/03_кейс_талай_каналы_итоги.md",
        ROOT / "docs/working/17_добор_ан_конкуренты_пон.md",
    ]
    for i, path in enumerate(parts):
        md = path.read_text(encoding="utf-8")
        md = md.replace("BD PMO", "внутренний рабочий ряд")
        md = md.replace("dual-use", "двойное использование")
        md = md.replace("Dual-use", "двойное использование")
        if i == 3:
            h = doc.add_paragraph()
            r = h.add_run(
                "Приложение. Добор эмпирической базы: посредники, конкуренты, общественные помещения"
            )
            set_run_font(r, size=14, bold=True)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(12)
            doc.add_page_break()
        mod.render_markdown(doc, md)
        if i < len(parts) - 1:
            doc.add_page_break()

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run(
        "Рынок жилья агломерации Горно-Алтайск — Майма · кейс «Талай» · IMPRO Group · 2026"
    )
    set_run_font(r, size=9, italic=True)

    doc.save(OUT)
    print(f"Saved {OUT} ({OUT.stat().st_size/1024:.1f} KB)")


if __name__ == "__main__":
    main()
